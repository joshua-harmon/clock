# Stars within 50 and 10 light years: lists, maps (positions, proper motions) and Word documents.
# Inputs in this folder:
#   naked_eye_stars_within_250ly.csv   (made by make_star_maps.py, Hipparcos 2007 + Hipparcos 1997 + IV/27A names)
#   ursa_major_stream_velocities.csv   (made by make_stream_map.py)
#   atlas_50ly.txt                     (R. Powell, An Atlas of the Universe, "Stars within 50 light years": Gliese numbers,
#                                       combined spectral classes of multiple stars, proper names, notes)
# Planets: confirmed planets per host from the NASA Exoplanet Archive (Planetary Systems Composite table), 11 Oct 2026.
# Stars fainter than Hipparcos within 10 ly: Gaia DR3 (ESA 2022).
# Order: python make_near_sets.py, then python make_plane_maps.py, then python make_near_sets.py again to put the plane maps in the documents.
import csv, math, os, re
import numpy as np
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

here = os.path.dirname(os.path.abspath(__file__))
T = np.array([[-0.0548755604, -0.8734370902, -0.4838350155],
              [ 0.4941094279, -0.4448296300,  0.7469822445],
              [-0.8676661490, -0.1980763734,  0.4559837762]])
K = 4.740470446                 # km/s per (mas/yr / mas)
KMS_LY_PER_YR = 1 / 299792.458  # 1 km/s in light years per year
YEARS = 10000                   # length of the proper-motion arrows

PLANETS_RAW = ("10138:1 102401:2 102409:4 106440:1 108870:1 109388:2 110443:1 11048:1 113020:4 114046:4 114622:6 116132:1 116727:1 116745:1 120005:1 "
 "14101:2 1475:2 15510:4 16537:1 16711:1 19394:3 21932:1 22627:1 22762:3 24186:1 27887:5 29295:2 3093:1 31300:1 31635:2 33226:2 36208:2 37826:1 40693:3 "
 "43587:5 45908:1 47103:3 47780:3 4856:1 48714:1 4872:1 49081:1 49189:1 51317:1 53721:3 53985:1 54035:2 54646:2 55042:1 5643:3 56528:3 56998:2 57050:2 "
 "57087:1 57548:1 59514:1 61706:1 62452:1 63510:1 64797:2 64924:3 65859:1 68469:2 70890:2 71253:1 74995:3 7513:3 75181:3 77257:1 77358:2 79431:1 79755:1 "
 "80337:1 80459:1 80824:3 8102:3 83043:1 85523:1 86057:1 86087:1 86162:2 86214:2 86287:1 87937:4 91768:1 91772:1 93101:1 94761:1 98698:1 99825:2")
PLANETS = {int(a): int(b) for a, b in (x.split(':') for x in PLANETS_RAW.split())}
CST = dict(And='Andromedae', Aql='Aquilae', Aqr='Aquarii', Ara='Arae', Ari='Arietis', Aur='Aurigae', Boo='Boötis', Cnc='Cancri', CVn='Canum Venaticorum',
 CMa='Canis Majoris', CMi='Canis Minoris', Cap='Capricorni', Car='Carinae', Cas='Cassiopeiae', Cen='Centauri', Cep='Cephei', Cet='Ceti', Col='Columbae',
 Com='Comae Berenices', CrB='Coronae Borealis', Crv='Corvi', Crt='Crateris', Cyg='Cygni', Del='Delphini', Dor='Doradus', Dra='Draconis', Eri='Eridani',
 For='Fornacis', Gem='Geminorum', Her='Herculis', Hya='Hydrae', Hyi='Hydri', Ind='Indi', Leo='Leonis', LMi='Leonis Minoris', Lep='Leporis', Lib='Librae',
 Lup='Lupi', Lyn='Lyncis', Lyr='Lyrae', Men='Mensae', Oph='Ophiuchi', Ori='Orionis', Pav='Pavonis', Peg='Pegasi', Per='Persei', Phe='Phoenicis',
 Psc='Piscium', PsA='Piscis Austrini', Pup='Puppis', Ret='Reticuli', Sco='Scorpii', Ser='Serpentis', Tau='Tauri', TrA='Trianguli Australis',
 Tri='Trianguli', Tuc='Tucanae', UMa='Ursae Majoris', UMi='Ursae Minoris', Vel='Velorum', Vir='Virginis')

def full_name(short):
    p = short.split(' ', 1)
    return p[0] + ' ' + CST[p[1]] if len(p) == 2 and p[1] in CST else short.replace('HR ', 'HR')

def gal(ra, dec):
    a, d = math.radians(ra), math.radians(dec)
    v = T @ np.array([math.cos(d)*math.cos(a), math.cos(d)*math.sin(a), math.sin(d)])
    return math.degrees(math.atan2(v[1], v[0])) % 360, math.degrees(math.asin(v[2]))

def radec_from_lb(l, b):
    lr, br = math.radians(l), math.radians(b)
    e = T.T @ np.array([math.cos(br)*math.cos(lr), math.cos(br)*math.sin(lr), math.sin(br)])
    return math.degrees(math.atan2(e[1], e[0])) % 360, math.degrees(math.asin(e[2]))

def state(ra, dec, plx, pmra, pmde, rv):
    """position (ly, galactic x y z) and velocity (ly per year, galactic) of a star"""
    a, d = math.radians(ra), math.radians(dec)
    re = np.array([math.cos(d)*math.cos(a), math.cos(d)*math.sin(a), math.sin(d)])
    p = np.array([-math.sin(a), math.cos(a), 0.0]); q = np.array([-math.sin(d)*math.cos(a), -math.sin(d)*math.sin(a), math.cos(d)])
    dist = 3261.56 / plx
    v = (K*pmra/plx)*p + (K*pmde/plx)*q + (rv or 0.0)*re
    return T @ (re*dist), T @ v * KMS_LY_PER_YR

def lb_of(x):
    r = np.linalg.norm(x); return math.degrees(math.atan2(x[1], x[0])) % 360, math.degrees(math.asin(x[2]/r)), r

hms = lambda ra: '%02d %04.1f' % (int(ra/15), (ra/15 - int(ra/15))*60)
def dms(d):
    s = '-' if d < 0 else '+'; d = abs(d); m = int(round((d - int(d))*60)); dd = int(d)
    if m == 60: dd, m = dd + 1, 0
    return '%s%02d %02d' % (s, dd, m)

# ------------------------------------------------------------------ data
uma, rvs = {}, {}
for r in csv.DictReader(open(os.path.join(here, 'ursa_major_stream_velocities.csv'), encoding='utf-8')):
    uma[int(r['HIP'])] = r['Stream match']; rvs[int(r['HIP'])] = float(r['Radial velocity (km/s)'])

def parse_atlas():
    out = []
    for line in open(os.path.join(here, 'atlas_50ly.txt'), encoding='utf-8'):
        if not re.match(r'^.{25}\d\d \d\d\.\d [+-]\d\d \d\d', line): continue
        m = re.match(r'(.{25})(\d\d) (\d\d\.\d) ([+-])(\d\d) (\d\d)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(G[lJ] *[\d.]+|--)', line)
        name = m.group(1).strip(); ra = (int(m.group(2)) + float(m.group(3))/60)*15
        dec = (1 if m.group(4) == '+' else -1)*(int(m.group(5)) + int(m.group(6))/60)
        rest = line.rstrip('\n')
        out.append(dict(name=name, ra=ra, dec=dec, cls=m.group(9), gl=re.sub(r'\s+', ' ', m.group(14)).replace('--', ''),
                        proper=rest[108:124].strip(), notes=rest[124:].strip()))
    return out
atlas = parse_atlas()

stars = []
for r in csv.DictReader(open(os.path.join(here, 'naked_eye_stars_within_250ly.csv'), encoding='utf-8')):
    l, b = float(r['Galactic l']), float(r['Galactic b']); ra, dec = radec_from_lb(l, b)
    stars.append(dict(HIP=int(r['HIP']), short=r['Name'], proper=r['Proper name'], ra=ra, dec=dec, l=l, b=b, sp=r['Spectral type'],
                      V=float(r['V mag']), plx=float(r['Parallax (mas)']), ly=float(r['Distance (ly)']),
                      pmra=float(r['pm RA (mas/yr)']), pmde=float(r['pm Dec (mas/yr)'])))

def sep(a, b):
    ra1, d1, ra2, d2 = map(math.radians, (a['ra'], a['dec'], b['ra'], b['dec']))
    c = math.sin(d1)*math.sin(d2) + math.cos(d1)*math.cos(d2)*math.cos(ra1 - ra2)
    return math.degrees(math.acos(max(-1, min(1, c))))

def planet_note(n):
    return '' if not n else ('Has a planet' if n == 1 else 'Has %d planets' % n)

def merge_notes(old, n_planets, extra):
    old = re.sub(r'\.?\s*Has (a|\d+) planets?\.?', '', old).strip(' .')
    parts = [x for x in (old, planet_note(n_planets), extra) if x]
    return '. '.join(parts)

# ------------------------------------------------------------------ 50 light years
near = [s for s in stars if s['ly'] <= 50.0]
used, rows50 = set(), []
for s in sorted(near, key=lambda s: s['V']):          # brighter components claim the Atlas entry first
    best = min(atlas, key=lambda a: sep(a, s)); d = sep(best, s)
    if d < 0.15 and best['name'] not in used:
        used.add(best['name']); s['atlas'] = best
    elif d < 0.15:
        s['companion_of'] = best['name']             # second Hipparcos entry of a double star already listed
rows50 = [s for s in near if 'companion_of' not in s]
for s in rows50:
    a = s.get('atlas')
    s['cname'] = a['name'] if a else full_name(s['short'])
    s['cls'] = a['cls'] if a else s['sp']
    s['gl'] = a['gl'] if a else ''
    s['pname'] = a['proper'] if a else s['proper']
    um = uma.get(s['HIP'], '')
    extra = {'nucleus': 'Ursa Major group nucleus', 'close (within 3 km/s)': 'Moves with the Ursa Major stream',
             'possible (3-6 km/s)': 'Possible Ursa Major stream member'}.get(um, '')
    if s['HIP'] == 71683: extra = (extra + '. ' if extra else '') + 'Proxima Centauri, 0.2 ly away, has 2 planets'
    s['notes'] = merge_notes(a['notes'] if a else '', PLANETS.get(s['HIP'], 0), extra)
    s['M'] = s['V'] + 5*math.log10(s['plx']) - 10
rows50.sort(key=lambda s: s['ly'])
missing = [a['name'] for a in atlas if a['name'] not in used]
print('50 ly:', len(rows50), 'systems;', len(near) - len(rows50), 'companions folded in; Atlas entries not within 50 ly now:', missing)
print('new since the Atlas list:', [s['cname'] for s in rows50 if 'atlas' not in s])

# ------------------------------------------------------------------ 10 light years
def hip(h): return next(s for s in stars if s['HIP'] == h)
ten = []
for h, name, cls, gl, pname, rv in [(71683, 'Alpha Centauri A', 'G2V', 'Gl 559 A', 'Rigil Kentaurus', -21.4),
                                    (71681, 'Alpha Centauri B', 'K1V', 'Gl 559 B', 'Toliman', -20.6),
                                    (32349, 'Alpha Canis Majoris A', 'A1V', 'Gl 244 A', 'Sirius', -5.5)]:
    s = dict(hip(h)); s.update(cname=name, cls=cls, gl=gl, pname=pname, rv=rv, src='Hipparcos 2007'); ten.append(s)
# Gaia DR3: ra, dec, parallax, pmra, pmdec, radial velocity, (V and type from the literature)
GAIA = [('Proxima Centauri', 'M5.5V', 'Gl 551', 'Proxima', 217.3923, -62.6761, 768.07, -3781.74, 769.47, -21.94, 11.13, 70890),
        ('Barnard’s Star', 'M4V', 'Gl 699', '', 269.4485, 4.7394, 546.98, -801.55, 10362.39, -110.47, 9.54, 87937),
        ('Wolf 359', 'M6V', 'Gl 406', '', 164.1032, 7.0027, 415.18, -3866.34, -2699.21, 19.5, 13.51, 0),
        ('Lalande 21185', 'M2V', 'Gl 411', '', 165.8310, 35.9487, 392.75, -580.06, -4776.59, -85.11, 7.52, 54035),
        ('Alpha Canis Majoris B', 'DA2', 'Gl 244 B', 'Sirius B', 101.2866, -16.7209, 379.21, -546.01, -1223.07, -5.5, 8.44, 0),   # Sirius B: system parallax and motion of Sirius A (Gaia's own values include B's orbit)
        ('Luyten 726-8 A', 'M5.5V', 'Gl 65 A', 'BL Ceti', 24.7716, -17.9483, 367.71, 3385.32, 544.39, 29.0, 12.54, 0),
        ('Luyten 726-8 B', 'M6V', 'Gl 65 B', 'UV Ceti', 24.7717, -17.9477, 373.84, 3178.69, 584.06, 32.0, 12.99, 0),
        ('Ross 154', 'M3.5V', 'Gl 729', '', 282.4588, -23.8371, 336.03, 639.37, -193.96, -11.12, 10.43, 0)]
for name, cls, gl, pname, ra, dec, plx, pmra, pmde, rv, V, h in GAIA:
    l, b = gal(ra, dec)
    ten.append(dict(HIP=h, cname=name, cls=cls, gl=gl, pname=pname, ra=ra, dec=dec, l=l, b=b, V=V, plx=plx, ly=3261.56/plx,
                    pmra=pmra, pmde=pmde, rv=rv, src='Gaia DR3'))
for s in ten:
    s['M'] = s['V'] + 5*math.log10(s['plx']) - 10
    s['pm'] = math.hypot(s['pmra'], s['pmde']) / 1000
    s['vt'] = K*s['pm']*1000/s['plx']
    n = PLANETS.get(s['HIP'], 0)
    s['notes'] = planet_note(n)
ten.sort(key=lambda s: s['ly'])
LAB = {'Alpha Centauri A': 'Alpha Centauri A, B and Proxima', 'Alpha Centauri B': '', 'Proxima Centauri': '',
       'Alpha Canis Majoris A': 'Sirius A and B', 'Alpha Canis Majoris B': '', 'Luyten 726-8 A': 'BL and UV Ceti (Luyten 726-8)', 'Luyten 726-8 B': ''}
for s_ in ten: s_['lab'] = LAB.get(s_['cname'], s_['pname'] or s_['cname'])
notes10 = {'Proxima Centauri': 'Nearest star to the Sun, 0.2 ly from Alpha Centauri A and B',
           'Barnard’s Star': 'Fastest proper motion of any star, 10.4″ a year: the width of the full Moon in 180 years',
           'Alpha Canis Majoris B': 'White dwarf companion of Sirius', 'Luyten 726-8 B': 'UV Ceti, the prototype flare star',
           'Wolf 359': 'Flare star', 'Lalande 21185': 'Brightest red dwarf in the northern sky (still too faint for the naked eye)'}
for s in ten:
    s['notes'] = '. '.join(x for x in (notes10.get(s['cname'], ''), s['notes']) if x)

# ------------------------------------------------------------------ figures
X = lambda l: math.radians(-(((l + 180) % 360) - 180))
TICKS = ['150°', '120°', '90°', '60°', '30°', '0°', '330°', '300°', '270°', '240°', '210°']

def sky_map(rows, maxly, fname, title, label_all=False):
    fig = plt.figure(figsize=(14, 7.6), dpi=130); ax = fig.add_subplot(111, projection='aitoff')
    sc = ax.scatter([X(r['l']) for r in rows], [math.radians(r['b']) for r in rows], s=[max(10, (6.6 - min(r['V'], 6.5))**2*4) for r in rows],
                    c=[r['ly'] for r in rows], cmap='viridis_r', vmin=0, vmax=maxly, edgecolors='none', zorder=3)
    P = [r for r in rows if PLANETS.get(r['HIP'], 0)]
    ax.scatter([X(r['l']) for r in P], [math.radians(r['b']) for r in P], s=150, facecolors='none', edgecolors='#d0302a', lw=1.3, zorder=4,
               label='has known planets (%d)' % len(P))
    for r in rows:
        nm = r['lab'] if 'lab' in r else (r['pname'] or (r['cname'] if label_all or r['cname'] == 'Zeta Reticuli' else ''))
        if nm: ax.annotate(nm, (X(r['l']), math.radians(r['b'])), xytext=(5, 4), textcoords='offset points', fontsize=7.5)
    ax.set_xticklabels(TICKS, fontsize=8); ax.grid(True, color='#bbb', lw=.5)
    cb = fig.colorbar(sc, ax=ax, orientation='horizontal', fraction=.04, pad=.07, shrink=.5); cb.set_label('distance (light years)')
    ax.set_title(title, fontsize=13, pad=14); ax.legend(loc='lower right', fontsize=8, frameon=False)
    fig.text(.5, .015, 'Galactic centre in the middle, longitude increasing to the left. Dot size: brightness. Planets: NASA Exoplanet Archive.', ha='center', fontsize=8, color='#555')
    fig.savefig(os.path.join(here, fname), bbox_inches='tight'); plt.close(fig)

def motion_map(rows, years, fname, title, label_all=False):
    """sky map with arrows from today's position to the position after `years` (full 3-D straight-line motion)"""
    fig = plt.figure(figsize=(14, 7.6), dpi=130); ax = fig.add_subplot(111, projection='aitoff')
    for r in rows:
        x, v = state(r['ra'], r['dec'], r['plx'], r['pmra'], r['pmde'], r.get('rv', rvs.get(r['HIP'], 0.0)))
        l1, b1, _ = lb_of(x + v*years)
        x0, y0, x1, y1 = X(r['l']), math.radians(r['b']), X(l1), math.radians(b1)
        if abs(x1 - x0) > math.pi: continue            # crosses the map edge: skip the arrow
        ax.annotate('', xy=(x1, y1), xytext=(x0, y0), arrowprops=dict(arrowstyle='-|>', color='#b0482c', lw=.8, mutation_scale=7), zorder=2)
    ax.scatter([X(r['l']) for r in rows], [math.radians(r['b']) for r in rows], s=[max(8, (6.6 - min(r['V'], 6.5))**2*3) for r in rows],
               c='#203a60', edgecolors='none', zorder=3)
    for r in rows:
        nm = r['lab'] if 'lab' in r else (r['pname'] or (r['cname'] if label_all or r['cname'] == 'Zeta Reticuli' else ''))
        if nm: ax.annotate(nm, (X(r['l']), math.radians(r['b'])), xytext=(5, 4), textcoords='offset points', fontsize=7.5)
    ax.set_xticklabels(TICKS, fontsize=8); ax.grid(True, color='#bbb', lw=.5)
    ax.set_title(title, fontsize=13, pad=14)
    fig.text(.5, .03, 'Each arrow runs from the star’s place today to its place %s years from now (straight-line motion through space, including its speed toward or away from us).' % format(years, ','),
             ha='center', fontsize=8.5, color='#555')
    fig.savefig(os.path.join(here, fname), bbox_inches='tight'); plt.close(fig)

def positions(rows, maxly, fname, title, label_all=False):
    def xyz(r):
        lr, br = math.radians(r['l']), math.radians(r['b']); return r['ly']*math.cos(br)*math.cos(lr), r['ly']*math.cos(br)*math.sin(lr), r['ly']*math.sin(br)
    Q = np.array([xyz(r) for r in rows]); pl = np.array([bool(PLANETS.get(r['HIP'], 0)) for r in rows]); Vm = np.array([min(r['V'], 6.5) for r in rows])
    fig, axs = plt.subplots(1, 2, figsize=(15, 7.6), dpi=130)
    for ax, (i, j, lab) in zip(axs, [(0, 1, ('X: toward galactic centre (ly)', 'Y: direction of the Sun’s orbit (ly)')), (0, 2, ('X: toward galactic centre (ly)', 'Z: toward north galactic pole (ly)'))]):
        ax.scatter(Q[:, i], Q[:, j], s=np.clip((6.6 - Vm)**2*3, 8, None), c='#3b6fb0', alpha=.8, edgecolors='none')
        ax.scatter(Q[pl, i], Q[pl, j], s=90, facecolors='none', edgecolors='#d0302a', lw=1.2, label='has known planets')
        for R_ in (maxly/2, maxly):
            t = np.linspace(0, 2*np.pi, 300); ax.plot(R_*np.cos(t), R_*np.sin(t), color='#999', lw=.8, ls='--'); ax.text(R_*.72, R_*.72, '%g ly' % R_, color='#777', fontsize=8)
        ax.plot(0, 0, marker='*', color='#e8a400', ms=13, zorder=5); ax.text(maxly*.02, -maxly*.06, 'Sun', fontsize=9)
        seen = set()
        for k, r in enumerate(rows):
            nm = r['lab'] if 'lab' in r else (r['pname'] or (r['cname'] if label_all or r['cname'] == 'Zeta Reticuli' else ''))
            if nm and nm not in seen and (label_all or r['V'] < 2.3 or nm in ('Sirius', 'Procyon', 'Zeta Reticuli')):
                seen.add(nm); ax.annotate(nm, (Q[k, i], Q[k, j]), xytext=(4, 3), textcoords='offset points', fontsize=7.5)
        m = maxly*1.06; ax.set_xlim(-m, m); ax.set_ylim(-m, m); ax.set_aspect('equal'); ax.set_xlabel(lab[0]); ax.set_ylabel(lab[1]); ax.grid(color='#eee')
    axs[0].set_title('Top view: the galactic plane seen from the north'); axs[1].set_title('Side view: up and down out of the plane'); axs[0].legend(loc='lower left', fontsize=8, frameon=False)
    fig.suptitle(title, fontsize=13); fig.savefig(os.path.join(here, fname), bbox_inches='tight'); plt.close(fig)

def approach_chart(rows, fname):
    """distance from the Sun over time, -40,000 to +60,000 years, straight-line motion"""
    t = np.linspace(-40000, 60000, 1001)
    fig, ax = plt.subplots(figsize=(13, 7), dpi=130)
    out = []
    for r in rows:
        x, v = state(r['ra'], r['dec'], r['plx'], r['pmra'], r['pmde'], r['rv'])
        d = np.linalg.norm(x[None, :] + t[:, None]*v[None, :], axis=1)
        k = int(np.argmin(d)); out.append((r['cname'], t[k], d[k]))
        col = {'Alpha Centauri A': '#1f77b4', 'Alpha Centauri B': '#1f77b4', 'Proxima Centauri': '#1f77b4', 'Barnard’s Star': '#d62728', 'Wolf 359': '#9467bd',
               'Lalande 21185': '#8c564b', 'Alpha Canis Majoris A': '#17becf', 'Alpha Canis Majoris B': '#17becf', 'Luyten 726-8 A': '#bcbd22',
               'Luyten 726-8 B': '#bcbd22', 'Ross 154': '#2ca02c'}.get(r['cname'], '#333')
        ls = {'Alpha Centauri B': '--', 'Proxima Centauri': ':', 'Alpha Canis Majoris B': '--', 'Luyten 726-8 B': '--'}.get(r['cname'], '-')
        ax.plot(t, d, lw=1.5, color=col, ls=ls, label=r['cname'] + (' (' + r['pname'] + ')' if r['pname'] and r['pname'] not in r['cname'] else ''))
        ax.plot(t[k], d[k], 'o', ms=4, color=ax.lines[-1].get_color())
    ax.axvline(0, color='#888', lw=.8); ax.text(500, ax.get_ylim()[1]*.97 if False else 19.2, 'today', fontsize=9, color='#555')
    ax.set_xlabel('years from now'); ax.set_ylabel('distance from the Sun (light years)'); ax.set_ylim(0, 20); ax.grid(color='#eee')
    ax.set_title('The nearest stars: distance from the Sun over 100,000 years (dots: closest approach)', fontsize=12.5)
    ax.legend(fontsize=8, ncol=2, frameon=False, loc='upper right')
    fig.text(.5, .005, 'Straight-line motion from today’s positions, proper motions and radial velocities (Hipparcos 2007, Gaia DR3); the Galaxy’s pull bends the paths slightly over these times.', ha='center', fontsize=8, color='#555')
    fig.savefig(os.path.join(here, fname), bbox_inches='tight'); plt.close(fig)
    return out

sky_map(rows50, 50, 'sky_galactic_within_50ly.png', 'Naked-eye stars within 50 light years (%d systems), galactic coordinates' % len(rows50))
positions(rows50, 50, 'positions_around_sun_50ly.png', 'Naked-eye stars within 50 light years around the Sun (%d systems)' % len(rows50))
motion_map(rows50, YEARS, 'proper_motion_within_50ly.png', 'Where the stars within 50 light years are heading: motion over the next %s years' % format(YEARS, ','))
sky_map(ten, 10, 'sky_galactic_within_10ly.png', 'All known stars within 10 light years (%d), galactic coordinates' % len(ten), label_all=True)
positions(ten, 10, 'positions_around_sun_10ly.png', 'All known stars within 10 light years around the Sun (%d)' % len(ten), label_all=True)
motion_map(ten, YEARS, 'proper_motion_within_10ly.png', 'Where the stars within 10 light years are heading: motion over the next %s years' % format(YEARS, ','), label_all=True)
closest = approach_chart(ten, 'nearest_stars_distance_over_time.png')
for n, t, d in closest: print('  closest approach: %-24s %+8.0f yr  %.2f ly' % (n, t, d))

# ------------------------------------------------------------------ CSV lists
COLS = ['Catalog Name', 'RA (h m)', 'Dec (d m)', 'Galactic l', 'Galactic b', 'Classification', 'Visual Mag', 'Abs Mag', 'Prllx (mas)', 'Dist (ly)',
        'Gliese Number', 'Proper Name', 'Notes']
def line(s):
    return [s['cname'], hms(s['ra']), dms(s['dec']), '%.1f' % s['l'], '%+.1f' % s['b'], s['cls'], '%.2f' % s['V'], '%.2f' % s['M'],
            '%.2f' % s['plx'], '%.2f' % s['ly'], s['gl'], s['pname'], s['notes']]
SUN = ['Sun', '---', '---', '--', '--', 'G2V', '-26.74', '4.83', '--', '0.00', '--', 'Sol', 'Has 8 planets']
with open(os.path.join(here, 'stars_within_50ly.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f); w.writerow(COLS + ['HIP']); w.writerow(SUN + [''])
    for s in rows50: w.writerow(line(s) + [s['HIP']])
COLS10 = COLS[:-1] + ['Proper motion (″/yr)', 'Tangential speed (km/s)', 'Radial velocity (km/s)', 'Notes']
def line10(s): return line(s)[:-1] + ['%.3f' % s['pm'], '%.1f' % s['vt'], '%+.1f' % s['rv'], s['notes']]
with open(os.path.join(here, 'stars_within_10ly.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f); w.writerow(COLS10); w.writerow(SUN[:-1] + ['--', '--', '--', SUN[-1]])
    for s in ten: w.writerow(line10(s))

# ------------------------------------------------------------------ Word documents
def new_doc(title):
    doc = Document(); sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.3))
    doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10)
    doc.add_heading(title, 0); return doc

def table(doc, cols, data, widths, size=7.5):
    t = doc.add_table(rows=1, cols=len(cols)); t.style = 'Light Grid Accent 1'
    for c, txt in zip(t.rows[0].cells, cols):
        c.text = ''; r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.size = Pt(size)
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    for vals in data:
        cells = t.add_row().cells
        for c, v in zip(cells, vals):
            c.text = ''; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, wd in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(wd*567)))
    for col, wd in zip(t.columns, widths): col.width = Cm(wd)
    for row in t.rows:
        for c, wd in zip(row.cells, widths): c.width = Cm(wd)

COLNOTE = ('Column 1: Bayer/Flamsteed designation, or the star’s number in the Harvard Revised Bright Star Catalogue (HR). Columns 2–3: right ascension (hours, minutes) '
           'and declination (degrees, minutes), epoch 2000. Columns 4–5: galactic longitude and latitude. Column 6: spectral class of the stars in the system. '
           'Column 7: apparent visual magnitude. Column 8: absolute magnitude (the magnitude at 32.6 light years). Column 9: parallax in milliarcseconds. '
           'Column 10: distance in light years (= 3261.56 / parallax). Column 11: number in the Gliese Catalogue of Nearby Stars. Column 12: proper name. Column 13: notes.')
SOURCES = ['van Leeuwen F. (2007), Hipparcos, the New Reduction of the Raw Data (parallaxes, proper motions), VizieR I/311.',
           'ESA (1997), The Hipparcos and Tycho Catalogues (magnitudes, spectral types), VizieR I/239.',
           'Kostjuk N.D. (2002), HD-DM-GC-HR-HIP-Bayer-Flamsteed Cross Index, VizieR IV/27A.',
           'Gaia Collaboration (2023), Gaia Data Release 3 (stars too faint for Hipparcos), ESA Gaia Archive.',
           'Gontcharov G.A. (2006), Pulkovo Compilation of Radial Velocities for 35,495 Hipparcos stars, VizieR III/252.',
           'NASA Exoplanet Archive, Planetary Systems Composite Parameters, exoplanetarchive.ipac.caltech.edu, accessed 11 October 2026.',
           'Gliese W., Jahreiß H. (1991), Third Catalogue of Nearby Stars; Hoffleit D., Warren W. (1991), Bright Star Catalogue, 5th ed.',
           'King J.R. et al. (2003), Stellar Kinematic Groups II: the Ursa Major Group, Astronomical Journal 125, 1980.',
           'Powell R., An Atlas of the Universe, “Stars within 50 light years” (atlasoftheuniverse.com): layout, Gliese numbers, combined spectral classes and notes.']

doc = new_doc('Stars within 50 Light Years')
doc.add_paragraph('Every star system within 50 light years that can be seen with the naked eye (magnitude 6.5 or brighter): %d systems, plus the Sun. '
                  'The list follows the Atlas of the Universe table with the distances brought up to date from the 2007 Hipparcos reduction, and the planets '
                  'brought up to date from the NASA Exoplanet Archive (11 October 2026): %d of these systems have confirmed planets. Notes also mark the stars that move '
                  'with the Ursa Major stream.' % (len(rows50), sum(1 for s in rows50 if PLANETS.get(s['HIP'], 0))))
table(doc, COLS, [SUN] + [line(s) for s in rows50], [3.4, 1.35, 1.35, 1.1, 1.1, 2.9, 1.1, 1.1, 1.35, 1.2, 1.4, 2.2, 5.0])
doc.add_paragraph(COLNOTE)
if missing: doc.add_paragraph('On the Atlas list but now measured just beyond 50 light years: ' + ', '.join(missing) + '.')
doc.add_heading('Maps', 1)
for f in (('plane_map_within_50ly.png',) if os.path.exists(os.path.join(here, 'plane_map_within_50ly.png')) else ()) + ('sky_galactic_within_50ly.png', 'positions_around_sun_50ly.png', 'proper_motion_within_50ly.png'):
    doc.add_picture(os.path.join(here, f), width=Cm(24.5))
doc.add_heading('Sources', 1)
for s in SOURCES: doc.add_paragraph(s, style='List Bullet')
doc.save(os.path.join(here, 'Stars within 50 Light Years.docx'))

doc = new_doc('Stars within 10 Light Years')
doc.add_paragraph('Every known star within 10 light years: %d stars in 7 systems, plus the Sun. Only Sirius and Alpha Centauri A and B are bright enough '
                  'to see with the naked eye; the rest are red dwarfs and one white dwarf. Bright stars from Hipparcos (2007), faint ones from Gaia DR3. '
                  'Two brown dwarf systems also lie within 10 light years, Luhman 16 A and B (6.5 ly) and WISE 0855−0714 (7.4 ly); they are not stars '
                  'and are too faint for both catalogues, so they are not listed. Planets: NASA Exoplanet Archive, 11 October 2026.' % len(ten))
table(doc, COLS10, [SUN[:-1] + ['--', '--', '--', SUN[-1]]] + [line10(s) for s in ten],
      [3.0, 1.3, 1.3, 1.0, 1.0, 1.4, 1.0, 1.0, 1.3, 1.1, 1.4, 1.8, 1.4, 1.4, 1.4, 4.5], size=7.5)
doc.add_paragraph(COLNOTE.replace('Column 13: notes.', 'Proper motion: the star’s motion across the sky in arcseconds per year; tangential speed: the same motion '
                  'in km/s at the star’s distance; radial velocity: speed toward (−) or away from (+) the Sun.'))
doc.add_paragraph('Closest approach to the Sun (straight-line motion): ' + '; '.join(('%s: still coming closer after 60,000 years (%.2f ly then)' % (n, d)) if t >= 59999 else ('%s %s years %s, at %.2f ly' % (n, format(abs(int(round(t, -2))), ','), 'from now' if t > 0 else 'ago', d)) for n, t, d in closest) + '.')
doc.add_heading('Maps', 1)
for f in (('plane_map_within_10ly.png',) if os.path.exists(os.path.join(here, 'plane_map_within_10ly.png')) else ()) + ('sky_galactic_within_10ly.png', 'positions_around_sun_10ly.png', 'proper_motion_within_10ly.png', 'nearest_stars_distance_over_time.png'):
    doc.add_picture(os.path.join(here, f), width=Cm(24.5))
doc.add_heading('Sources', 1)
for s in SOURCES[:6] + SOURCES[6:7]: doc.add_paragraph(s, style='List Bullet')
doc.save(os.path.join(here, 'Stars within 10 Light Years.docx'))
print('done')
